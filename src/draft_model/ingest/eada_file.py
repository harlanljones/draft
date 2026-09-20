"""First-party EADA ingest from the U.S. Department of Education.

Downloads the official EADA survey archive (U.S. federal public domain)
directly from ope.ed.gov, extracts the men's-baseball rows from the
per-institution sport file, and derives the same program-resource context
features the packaged ncaa_bbStats extract carries — but for every program
that files EADA, not only the package's bundled subset.

The result is a local CSV plus a provenance manifest under
``data/empirical/`` (gitignored). ``draft_model.ingest.eada`` prefers this
first-party index over the packaged extract when it exists.

Schemas: EADA workbooks before 2021-22 use per-sport suffixed columns
(``EXP_MEN_Baseball``); 2021-22 onward use generic per-sport columns in a
sport-row file (``EXP_MEN`` with ``Sports == "Baseball"``). Both are handled.
"""

from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from datetime import date
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

EADA_API_BASE = "https://ope.ed.gov/athletics/api/dataFiles"

# Generic (2021-22+) column → legacy suffixed column consumed by
# ncaa_bbStats.program_store.derive_features.
_GENERIC_TO_LEGACY = {
    "PARTIC_MEN": "PARTIC_MEN_Baseball",
    "OPEXPPERPART_MEN": "OPEXPPERPART_MEN_Baseball",
    "EXP_MEN": "EXP_MEN_Baseball",
    "REV_MEN": "REV_MEN_Baseball",
    "MEN_TOTAL_HEADCOACH": "MEN_TOTAL_HEADCOACH_Baseball",
    "MEN_TOTAL_ASSTCOACH": "MEN_TOTAL_ASSTCOACH_Baseball",
    "RECRUITEXP_MEN": "RECRUITEXP_MEN",
    "HDCOACH_SAL_FTE_MEN": "HDCOACH_SAL_FTE_MEN",
}


def eada_file_url(academic_end_year: int) -> str:
    """URL of the official EADA archive whose academic year ends in ``end_year``."""
    listing = _get_json(f"{EADA_API_BASE}/fileList")
    for entry in listing:
        # Match on the survey year and the primary per-institution archive;
        # filenames use spaces before 2021-22 and underscores after.
        link = str(entry.get("LinkName", ""))
        if entry.get("Year") == academic_end_year and link.startswith("Data for academic year"):
            return f"{EADA_API_BASE}/file?fileName={quote(str(entry['FileName']))}"
    raise ValueError(f"no official EADA file for academic year ending {academic_end_year}")


def download_eada_file(academic_end_year: int, dest_dir: Path) -> dict[str, str]:
    """Download (idempotently) the official EADA zip and record provenance."""
    dest_dir.mkdir(parents=True, exist_ok=True)
    url = eada_file_url(academic_end_year)
    file_name = url.split("fileName=")[-1]
    zip_path = dest_dir / file_name
    if not zip_path.exists():
        request = Request(url, headers={"User-Agent": "draft-model/0.1"})
        with urlopen(request, timeout=300) as response:  # noqa: S310 - fixed https host
            zip_path.write_bytes(response.read())
    return {
        "file": zip_path.name,
        "sha256": hashlib.sha256(zip_path.read_bytes()).hexdigest(),
        "source_url": eada_file_url(academic_end_year),
        "license_name": "Public domain (U.S. federal government work)",
        "retrieved_at": date.today().isoformat(),
    }


def extract_schools_workbook(zip_path: Path, work_dir: Path) -> Path:
    """Extract the per-institution sport workbook from an EADA zip."""
    work_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as archive:
        candidates = [n for n in archive.namelist() if n.lower() in ("schools.xlsx", "eada.xlsx")]
        if not candidates:
            raise ValueError(f"no schools workbook found in {zip_path.name}")
        return Path(archive.extract(candidates[0], work_dir))


def _get_json(url: str) -> list[dict[str, object]]:
    request = Request(url, headers={"User-Agent": "draft-model/0.1"})
    with urlopen(request, timeout=60) as response:  # noqa: S310 - fixed https host
        payload = json.loads(response.read().decode("utf-8"))
    return payload if isinstance(payload, list) else []


def normalize_rows(workbook_path: Path, eada_year: int) -> list[dict[str, object]]:
    """Return legacy-schema records for Baseball rows of one EADA workbook.

    Handles both the legacy suffixed-column schema and the 2021-22+
    generic-column sport-row schema. Columns the newer schema does not
    publish (recruiting expense, head-coach salary) are simply absent and
    derive_features treats them as missing.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    index = {name: i for i, name in enumerate(header)}

    if any(name in index for name in _GENERIC_TO_LEGACY.values()):
        legacy = True
        sport_col: str | None = None
    else:
        legacy = False
        sport_col = None
        for candidate in ("Sports", "SPORTSCODE"):
            if candidate in index:
                sport_col = candidate
                break
        if sport_col is None:
            raise ValueError(f"{workbook_path.name} has neither legacy nor sport-row schema")

    records: list[dict[str, object]] = []
    for values in rows:
        if not legacy:
            sport = values[index[sport_col or "Sports"]]
            if not isinstance(sport, str) or sport.strip().lower() != "baseball":
                continue
        record: dict[str, object] = {"eada_year": eada_year}
        for generic, legacy_name in _GENERIC_TO_LEGACY.items():
            source = legacy_name if legacy else generic
            if source in index:
                record[legacy_name] = values[index[source]]
        for shared in ("unitid", "institution_name", "state_cd"):
            if shared in index:
                record[shared] = values[index[shared]]
        if record.get("unitid") is not None:
            records.append(record)
    return records


def build_context_rows(
    workbook_path: Path, eada_year: int
) -> tuple[list[dict[str, object]], list[str]]:
    """Derive context features for one EADA workbook and join registry team_ids.

    Returns (rows-with-team_id, unmatched-institution-names). Joining runs
    through ncaa_bbStats' registry (IPEDS unitid first, then exact/alias
    names) — never fuzzy matching.
    """
    from ncaa_bbStats.program_store import attach_team_ids, derive_features

    features = derive_features(normalize_rows(workbook_path, eada_year))
    matched, unmatched = attach_team_ids(features)
    for row in matched:
        row["season"] = eada_year
    return matched, unmatched


def write_context_csv(rows: list[dict[str, object]], output_path: Path) -> None:
    """Persist derived context rows for the first-party index loader."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "team_id",
        "institution_name",
        "unitid",
        "season",
        "eada_year",
        "budget_pct",
        "log_budget_per_player",
        "roster_size",
    ]
    with open(output_path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def load_local_index(derived_dir: Path) -> list[dict[str, str]]:
    """Rows from any first-party eada-context CSV under ``derived_dir``."""
    rows: list[dict[str, str]] = []
    if not derived_dir.is_dir():
        return rows
    for path in sorted(derived_dir.glob("eada-context*.csv")):
        with open(path, encoding="utf-8") as handle:
            rows.extend(csv.DictReader(handle))
    return rows
